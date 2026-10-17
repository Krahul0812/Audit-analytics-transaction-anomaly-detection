import mysql.connector
import pandas as pd
from getpass import getpass
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import PieChart, BarChart, Reference
from openpyxl.utils import get_column_letter


# ============================================================
# AUTOMATED AUDIT ANALYTICS SYSTEM
# ============================================================
#
# INPUT:
#     Transaction.csv
#
# PROCESS:
#     CSV → MySQL → SQL Audit Checks → Python Risk Scoring
#
# OUTPUT:
#     Audit_Exception_Report.xlsx
#
# ============================================================


CSV_FILE = "Transaction.csv"
OUTPUT_FILE = "Audit_Exception_Report.xlsx"

MYSQL_HOST = "localhost"
MYSQL_USER = "root"
MYSQL_DATABASE = "audit_analytics"


# ============================================================
# 1. GET MYSQL PASSWORD
# ============================================================

print("=" * 70)
print("          AUTOMATED AUDIT ANALYTICS SYSTEM")
print("=" * 70)

MYSQL_PASSWORD = getpass(
    "\nEnter MySQL password: "
)


# ============================================================
# 2. LOAD CSV
# ============================================================

print("\n[1/7] Reading transaction data...")

data = pd.read_csv(CSV_FILE)

print("Transactions loaded:", len(data))


# ============================================================
# 3. VALIDATE REQUIRED COLUMNS
# ============================================================

required_columns = [
    "Transaction_ID",
    "Vendor_Name",
    "Invoice_Number",
    "Invoice_Date",
    "Payment_Date",
    "Amount",
    "GST_Rate",
    "Approval_Status",
    "Department",
    "Payment_Time"
]

missing_columns = [
    column
    for column in required_columns
    if column not in data.columns
]

if missing_columns:

    print("\nERROR: Missing columns:")

    for column in missing_columns:
        print("-", column)

    raise SystemExit(
        "\nPlease check the CSV format."
    )


# ============================================================
# 4. CLEAN DATA
# ============================================================

print("\n[2/7] Cleaning transaction data...")

data["Invoice_Date"] = pd.to_datetime(
    data["Invoice_Date"],
    errors="coerce"
)

data["Payment_Date"] = pd.to_datetime(
    data["Payment_Date"],
    errors="coerce"
)

data["Amount"] = pd.to_numeric(
    data["Amount"],
    errors="coerce"
)

data["GST_Rate"] = pd.to_numeric(
    data["GST_Rate"],
    errors="coerce"
)

# Remove completely invalid rows
data = data.dropna(
    subset=[
        "Transaction_ID",
        "Invoice_Date",
        "Payment_Date",
        "Amount"
    ]
)

print(
    "Valid transactions:",
    len(data)
)


# ============================================================
# 5. CONNECT TO MYSQL
# ============================================================

print("\n[3/7] Connecting to MySQL...")

connection = mysql.connector.connect(
    host=MYSQL_HOST,
    user=MYSQL_USER,
    password=MYSQL_PASSWORD,
    database=MYSQL_DATABASE
)

print("MySQL connection successful!")


# ============================================================
# 6. REPLACE OLD TRANSACTION DATA
# ============================================================

print("\n[4/7] Loading transactions into MySQL...")

cursor = connection.cursor()

# Remove old data
cursor.execute(
    "TRUNCATE TABLE transactions"
)

insert_query = """
INSERT INTO transactions
(
    Transaction_ID,
    Vendor_Name,
    Invoice_Number,
    Invoice_Date,
    Payment_Date,
    Amount,
    GST_Rate,
    Approval_Status,
    Department,
    Payment_Time
)
VALUES
(
    %s, %s, %s, %s, %s,
    %s, %s, %s, %s, %s
)
"""

for _, row in data.iterrows():

    cursor.execute(
        insert_query,
        (
            row["Transaction_ID"],
            row["Vendor_Name"],
            row["Invoice_Number"],
            row["Invoice_Date"].date(),
            row["Payment_Date"].date(),
            float(row["Amount"]),
            float(row["GST_Rate"]),
            row["Approval_Status"],
            row["Department"],
            str(row["Payment_Time"])
        )
    )

connection.commit()

print(
    "Transactions loaded into MySQL:",
    len(data)
)


# ============================================================
# 7. SQL AUDIT QUERY
# ============================================================

print("\n[5/7] Running SQL audit checks...")

audit_query = """
SELECT
    Transaction_ID,
    Vendor_Name,
    Invoice_Number,
    Invoice_Date,
    Payment_Date,
    Amount,
    GST_Rate,
    Approval_Status,
    Department,
    Payment_Time
FROM transactions

WHERE

    Payment_Date < Invoice_Date

    OR Amount > 100000

    OR LOWER(Approval_Status) = 'pending'

    OR Payment_Time < '06:00:00'

    OR Payment_Time >= '22:00:00'

    OR Invoice_Number IN
    (
        SELECT Invoice_Number
        FROM transactions
        GROUP BY Invoice_Number
        HAVING COUNT(*) > 1
    )
"""


cursor.execute(audit_query)

rows = cursor.fetchall()

columns = [
    column[0]
    for column in cursor.description
]

audit_data = pd.DataFrame(
    rows,
    columns=columns
)

print(
    "Potential exceptions found:",
    len(audit_data)
)


# ============================================================
# 8. PYTHON RISK ANALYSIS
# ============================================================

print("\n[6/7] Calculating risk scores...")

audit_data["Risk_Score"] = 0

audit_data["Risk_Flags"] = ""


# ------------------------------------------------------------
# RULE 1: PAYMENT BEFORE INVOICE
# ------------------------------------------------------------

rule1 = (
    audit_data["Payment_Date"]
    <
    audit_data["Invoice_Date"]
)

audit_data.loc[
    rule1,
    "Risk_Score"
] += 30

audit_data.loc[
    rule1,
    "Risk_Flags"
] += "Payment before invoice; "


# ------------------------------------------------------------
# RULE 2: DUPLICATE INVOICE
# ------------------------------------------------------------

duplicate_invoices = (
    audit_data["Invoice_Number"]
    .duplicated(
        keep=False
    )
)

audit_data.loc[
    duplicate_invoices,
    "Risk_Score"
] += 30

audit_data.loc[
    duplicate_invoices,
    "Risk_Flags"
] += "Duplicate invoice; "


# ------------------------------------------------------------
# RULE 3: HIGH-VALUE TRANSACTION
# ------------------------------------------------------------

rule3 = (
    audit_data["Amount"]
    >
    100000
)

audit_data.loc[
    rule3,
    "Risk_Score"
] += 20

audit_data.loc[
    rule3,
    "Risk_Flags"
] += "High transaction amount; "


# ------------------------------------------------------------
# RULE 4: PENDING APPROVAL
# ------------------------------------------------------------

rule4 = (
    audit_data[
        "Approval_Status"
    ]
    .astype(str)
    .str.lower()
    ==
    "pending"
)

audit_data.loc[
    rule4,
    "Risk_Score"
] += 20

audit_data.loc[
    rule4,
    "Risk_Flags"
] += "Pending approval; "


# ------------------------------------------------------------
# RULE 5: UNUSUAL PAYMENT TIME
# ------------------------------------------------------------

def get_payment_hour(value):

    try:

        if hasattr(
            value,
            "total_seconds"
        ):

            seconds = (
                value.total_seconds()
            )

        else:

            time_value = pd.to_timedelta(
                str(value)
            )

            seconds = (
                time_value.total_seconds()
            )

        return seconds / 3600

    except:

        return 12


payment_hours = audit_data[
    "Payment_Time"
].apply(
    get_payment_hour
)

rule5 = (
    (payment_hours < 6)
    |
    (payment_hours >= 22)
)

audit_data.loc[
    rule5,
    "Risk_Score"
] += 15

audit_data.loc[
    rule5,
    "Risk_Flags"
] += "Unusual payment time; "


# ============================================================
# 9. RISK LEVEL
# ============================================================

def calculate_risk_level(score):

    if score >= 30:
        return "HIGH"

    elif score > 0:
        return "MEDIUM"

    else:
        return "LOW"


audit_data["Risk_Level"] = (
    audit_data[
        "Risk_Score"
    ]
    .apply(
        calculate_risk_level
    )
)


# ============================================================
# 10. AUDIT REMARKS
# ============================================================

audit_data["Audit_Remarks"] = (
    audit_data["Risk_Flags"]
    .apply(
        lambda x:
        "Review transaction and supporting documents."
    )
)


# ============================================================
# 11. CREATE EXCEL REPORT
# ============================================================

print("\n[7/7] Creating Excel report...")


# Convert date/time values for Excel
audit_data["Invoice_Date"] = pd.to_datetime(
    audit_data["Invoice_Date"]
)

audit_data["Payment_Date"] = pd.to_datetime(
    audit_data["Payment_Date"]
)


# Create workbook
workbook = Workbook()


# ------------------------------------------------------------
# DASHBOARD SHEET
# ------------------------------------------------------------

dashboard = workbook.active

dashboard.title = "Dashboard"


dashboard["A1"] = (
    "AUDIT ANALYTICS DASHBOARD"
)

dashboard["A1"].font = Font(
    bold=True,
    size=18
)

dashboard.merge_cells(
    "A1:F1"
)

dashboard["B3"] = (
    "Total Transactions"
)

dashboard["C3"] = (
    "Flagged Transactions"
)

dashboard["D3"] = (
    "High Risk"
)

dashboard["E3"] = (
    "Medium Risk"
)

dashboard["F3"] = (
    "Exception Rate"
)


total_transactions = len(data)

flagged_transactions = len(
    audit_data
)

high_risk = len(
    audit_data[
        audit_data["Risk_Level"]
        ==
        "HIGH"
    ]
)

medium_risk = len(
    audit_data[
        audit_data["Risk_Level"]
        ==
        "MEDIUM"
    ]
)

exception_rate = (
    flagged_transactions
    /
    total_transactions
    if total_transactions > 0
    else 0
)


dashboard["B4"] = total_transactions

dashboard["C4"] = flagged_transactions

dashboard["D4"] = high_risk

dashboard["E4"] = medium_risk

dashboard["F4"] = exception_rate

dashboard["F4"].number_format = "0%"


# KPI formatting
for cell in dashboard[3]:

    cell.font = Font(
        bold=True
    )

    cell.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )


for cell in dashboard[4]:

    cell.font = Font(
        bold=True,
        size=16
    )

    cell.alignment = Alignment(
        horizontal="center"
    )


# ============================================================
# RISK DATA
# ============================================================

dashboard["A7"] = "Risk Level"

dashboard["B7"] = "Count"

dashboard["A8"] = "HIGH"

dashboard["B8"] = high_risk

dashboard["A9"] = "MEDIUM"

dashboard["B9"] = medium_risk


# ============================================================
# RISK PIE CHART
# ============================================================

pie = PieChart()

labels = Reference(
    dashboard,
    min_col=1,
    min_row=8,
    max_row=9
)

values = Reference(
    dashboard,
    min_col=2,
    min_row=7,
    max_row=9
)

pie.add_data(
    values,
    titles_from_data=True
)

pie.set_categories(
    labels
)

pie.title = "Risk Distribution"

pie.height = 7

pie.width = 10

dashboard.add_chart(
    pie,
    "A11"
)


# ============================================================
# DEPARTMENT ANALYSIS
# ============================================================

department_counts = (
    audit_data[
        "Department"
    ]
    .value_counts()
)


dashboard["D7"] = (
    "Department"
)

dashboard["E7"] = (
    "Exceptions"
)


row_number = 8

for department, count in (
    department_counts.items()
):

    dashboard.cell(
        row=row_number,
        column=4,
        value=department
    )

    dashboard.cell(
        row=row_number,
        column=5,
        value=int(count)
    )

    row_number += 1


department_chart = BarChart()

department_chart.type = "col"

department_chart.title = (
    "Exceptions by Department"
)

department_chart.y_axis.title = (
    "Exceptions"
)

department_chart.x_axis.title = (
    "Department"
)


data_ref = Reference(
    dashboard,
    min_col=5,
    min_row=7,
    max_row=row_number - 1
)

category_ref = Reference(
    dashboard,
    min_col=4,
    min_row=8,
    max_row=row_number - 1
)

department_chart.add_data(
    data_ref,
    titles_from_data=True
)

department_chart.set_categories(
    category_ref
)

department_chart.height = 7

department_chart.width = 12

dashboard.add_chart(
    department_chart,
    "F7"
)


# ============================================================
# VENDOR ANALYSIS
# ============================================================

vendor_counts = (
    audit_data[
        "Vendor_Name"
    ]
    .value_counts()
)


dashboard["D15"] = (
    "Vendor"
)

dashboard["E15"] = (
    "Exceptions"
)


row_number = 16

for vendor, count in (
    vendor_counts.items()
):

    dashboard.cell(
        row=row_number,
        column=4,
        value=vendor
    )

    dashboard.cell(
        row=row_number,
        column=5,
        value=int(count)
    )

    row_number += 1


vendor_chart = BarChart()

vendor_chart.type = "col"

vendor_chart.title = (
    "Exceptions by Vendor"
)

vendor_chart.y_axis.title = (
    "Exceptions"
)

vendor_chart.x_axis.title = (
    "Vendor"
)


vendor_data = Reference(
    dashboard,
    min_col=5,
    min_row=15,
    max_row=row_number - 1
)

vendor_categories = Reference(
    dashboard,
    min_col=4,
    min_row=16,
    max_row=row_number - 1
)

vendor_chart.add_data(
    vendor_data,
    titles_from_data=True
)

vendor_chart.set_categories(
    vendor_categories
)

vendor_chart.height = 7

vendor_chart.width = 12

dashboard.add_chart(
    vendor_chart,
    "F22"
)


# ============================================================
# AUDIT EXCEPTIONS SHEET
# ============================================================

exceptions_sheet = workbook.create_sheet(
    "Audit Exceptions"
)

for column_number, column_name in enumerate(
    audit_data.columns,
    start=1
):

    cell = exceptions_sheet.cell(
        row=1,
        column=column_number,
        value=column_name
    )

    cell.font = Font(
        bold=True
    )

    cell.fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7"
    )


for row_number, row_data in enumerate(
    audit_data.itertuples(
        index=False,
        name=None
    ),
    start=2
):

    for column_number, value in enumerate(
        row_data,
        start=1
    ):

        exceptions_sheet.cell(
            row=row_number,
            column=column_number,
            value=value
        )


# ============================================================
# TRANSACTIONS SHEET
# ============================================================

transactions_sheet = workbook.create_sheet(
    "Transactions"
)

for column_number, column_name in enumerate(
    data.columns,
    start=1
):

    cell = transactions_sheet.cell(
        row=1,
        column=column_number,
        value=column_name
    )

    cell.font = Font(
        bold=True
    )


for row_number, row_data in enumerate(
    data.itertuples(
        index=False,
        name=None
    ),
    start=2
):

    for column_number, value in enumerate(
        row_data,
        start=1
    ):

        transactions_sheet.cell(
            row=row_number,
            column=column_number,
            value=value
        )


# ============================================================
# FORMAT ALL SHEETS
# ============================================================

for sheet in [
    dashboard,
    exceptions_sheet,
    transactions_sheet
]:

    sheet.freeze_panes = "A2"

    for column in sheet.columns:

        max_length = 0

        column_letter = (
            get_column_letter(
                column[0].column
            )
        )

        for cell in column:

            if cell.value is not None:

                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        sheet.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            35
        )


# ============================================================
# SAVE EXCEL REPORT
# ============================================================

workbook.save(
    OUTPUT_FILE
)


# ============================================================
# CLOSE MYSQL
# ============================================================

cursor.close()

connection.close()


# ============================================================
# FINAL SUMMARY
# ============================================================

print("\n" + "=" * 70)

print(
    "              AUDIT PROCESS COMPLETED"
)

print("=" * 70)

print(
    "\nTotal Transactions :",
    total_transactions
)

print(
    "Flagged Transactions:",
    flagged_transactions
)

print(
    "High Risk           :",
    high_risk
)

print(
    "Medium Risk         :",
    medium_risk
)

print(
    "Exception Rate      :",
    f"{exception_rate:.0%}"
)

print(
    "\nExcel report created:"
)

print(
    OUTPUT_FILE
)

print(
    "\nMySQL connection closed."
)

print(
    "\nYou can now replace Transaction.csv"
)

print(
    "and run this program again."
)