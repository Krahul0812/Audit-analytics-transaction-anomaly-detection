import mysql.connector
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ============================================================
# AUDIT ANALYTICS & TRANSACTION ANOMALY DETECTION SYSTEM
# MILESTONE 5 - AUTOMATED EXCEPTION REPORT
# ============================================================

# 1. CONNECT TO MYSQL
connection = mysql.connector.connect(
    host=" ",
    user=" ",
    password=" ",
    database=" "
)

print("=" * 70)
print("       AUDIT ANALYTICS - AUTOMATED EXCEPTION REPORT")
print("=" * 70)

# 2. SQL AUDIT QUERY
query = """
SELECT
    Transaction_ID,
    Vendor_Name,
    Invoice_Number,
    Invoice_Date,
    Payment_Date,
    Amount,
    Approval_Status,
    Department,
    Payment_Time
FROM transactions
WHERE Payment_Date < Invoice_Date
   OR Amount > 100000
   OR LOWER(Approval_Status) = 'pending'
   OR Payment_Time < '06:00:00'
   OR Payment_Time >= '22:00:00'
   OR Invoice_Number IN (
       SELECT Invoice_Number
       FROM transactions
       GROUP BY Invoice_Number
       HAVING COUNT(*) > 1
   );
"""

# 3. EXECUTE QUERY USING MYSQL CURSOR
cursor = connection.cursor()
cursor.execute(query)

rows = cursor.fetchall()
columns = [column[0] for column in cursor.description]

# Convert MySQL results into Pandas DataFrame
audit_data = pd.DataFrame(rows, columns=columns)

cursor.close()

# 4. CONVERT DATE COLUMNS
audit_data["Invoice_Date"] = pd.to_datetime(
    audit_data["Invoice_Date"]
)

audit_data["Payment_Date"] = pd.to_datetime(
    audit_data["Payment_Date"]
)

# 5. CREATE RISK SCORE AND FLAGS
audit_data["Risk_Score"] = 0
audit_data["Risk_Flags"] = ""

# ------------------------------------------------------------
# RULE 1: PAYMENT BEFORE INVOICE
# ------------------------------------------------------------

rule1 = audit_data["Payment_Date"] < audit_data["Invoice_Date"]

audit_data.loc[rule1, "Risk_Score"] += 30

audit_data.loc[rule1, "Risk_Flags"] += (
    "Payment before invoice; "
)

# ------------------------------------------------------------
# RULE 2: DUPLICATE INVOICE
# ------------------------------------------------------------

duplicate_invoices = audit_data["Invoice_Number"].duplicated(
    keep=False
)

audit_data.loc[duplicate_invoices, "Risk_Score"] += 30

audit_data.loc[duplicate_invoices, "Risk_Flags"] += (
    "Duplicate invoice; "
)

# ------------------------------------------------------------
# RULE 3: HIGH-VALUE TRANSACTION
# ------------------------------------------------------------

rule3 = audit_data["Amount"] > 100000

audit_data.loc[rule3, "Risk_Score"] += 20

audit_data.loc[rule3, "Risk_Flags"] += (
    "High transaction amount; "
)

# ------------------------------------------------------------
# RULE 4: PENDING APPROVAL
# ------------------------------------------------------------

rule4 = (
    audit_data["Approval_Status"]
    .astype(str)
    .str.lower()
    == "pending"
)

audit_data.loc[rule4, "Risk_Score"] += 20

audit_data.loc[rule4, "Risk_Flags"] += (
    "Pending approval; "
)

# ------------------------------------------------------------
# RULE 5: UNUSUAL PAYMENT TIME
# ------------------------------------------------------------

# MySQL TIME may be returned by Python as timedelta.
# Convert it safely into total seconds.

payment_seconds = audit_data["Payment_Time"].apply(
    lambda x: x.total_seconds()
    if hasattr(x, "total_seconds")
    else pd.to_timedelta(str(x)).total_seconds()
)

payment_hour = payment_seconds / 3600

rule5 = (
    (payment_hour < 6)
    | (payment_hour >= 22)
)

audit_data.loc[rule5, "Risk_Score"] += 15

audit_data.loc[rule5, "Risk_Flags"] += (
    "Unusual payment time; "
)

# ------------------------------------------------------------
# 6. CALCULATE RISK LEVEL
# ------------------------------------------------------------

def calculate_risk_level(score):

    if score >= 30:
        return "HIGH"

    elif score > 0:
        return "MEDIUM"

    else:
        return "LOW"


audit_data["Risk_Level"] = audit_data[
    "Risk_Score"
].apply(calculate_risk_level)

# ------------------------------------------------------------
# 7. AUDIT REMARKS
# ------------------------------------------------------------

audit_data["Audit_Remarks"] = audit_data[
    "Risk_Flags"
].apply(
    lambda x:
    "Review transaction and supporting documents."
)

# ------------------------------------------------------------
# 8. FINAL REPORT COLUMNS
# ------------------------------------------------------------

report = audit_data[
    [
        "Transaction_ID",
        "Vendor_Name",
        "Invoice_Number",
        "Invoice_Date",
        "Payment_Date",
        "Amount",
        "Approval_Status",
        "Department",
        "Payment_Time",
        "Risk_Score",
        "Risk_Level",
        "Risk_Flags",
        "Audit_Remarks"
    ]
]

# ------------------------------------------------------------
# 9. CREATE EXCEL REPORT
# ------------------------------------------------------------

output_file = "Audit_Exception_Report.xlsx"

report.to_excel(
    output_file,
    index=False,
    sheet_name="Audit Exceptions"
)

# ------------------------------------------------------------
# 10. FORMAT EXCEL REPORT
# ------------------------------------------------------------

workbook = load_workbook(output_file)

worksheet = workbook[
    "Audit Exceptions"
]

# Header formatting
for cell in worksheet[1]:

    cell.font = Font(
        bold=True
    )

    cell.fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7"
    )

    cell.alignment = Alignment(
        horizontal="center"
    )

# Amount formatting
amount_column = 6

for row in range(
    2,
    worksheet.max_row + 1
):

    worksheet.cell(
        row=row,
        column=amount_column
    ).number_format = '₹#,##0.00'

# General alignment
for row in worksheet.iter_rows():

    for cell in row:

        cell.alignment = Alignment(
            vertical="center",
            wrap_text=True
        )

# Adjust column widths
for column_cells in worksheet.columns:

    max_length = 0

    column_letter = get_column_letter(
        column_cells[0].column
    )

    for cell in column_cells:

        if cell.value is not None:

            max_length = max(
                max_length,
                len(str(cell.value))
            )

    worksheet.column_dimensions[
        column_letter
    ].width = min(
        max_length + 2,
        35
    )

# Freeze header
worksheet.freeze_panes = "A2"

# Save workbook
workbook.save(output_file)

# ------------------------------------------------------------
# 11. DISPLAY SUMMARY
# ------------------------------------------------------------

print("\n" + "=" * 70)
print("                 AUDIT REPORT SUMMARY")
print("=" * 70)

print(
    "Total Exceptions :",
    len(report)
)

print(
    "High Risk        :",
    len(
        report[
            report["Risk_Level"] == "HIGH"
        ]
    )
)

print(
    "Medium Risk      :",
    len(
        report[
            report["Risk_Level"] == "MEDIUM"
        ]
    )
)

print(
    "Low Risk         :",
    len(
        report[
            report["Risk_Level"] == "LOW"
        ]
    )
)

print("\nExcel report created successfully!")

print(
    "File:",
    output_file
)

# ------------------------------------------------------------
# 12. CLOSE MYSQL CONNECTION
# ------------------------------------------------------------

connection.close()

print("\nMySQL connection closed.")

print(
    "Audit process completed successfully."
)